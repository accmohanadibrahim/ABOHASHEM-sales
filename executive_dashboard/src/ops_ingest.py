"""
Operations data (inventory, purchases, suppliers, item-level returns).

Sources are copied unchanged into ``data/raw/ops/`` from the company's
operations repository (Abohashem):

  * ``inventory_valuation_2026-06.xlsx`` — June stock count with prices
    (raw materials block + finished-goods block).
  * ``inventory_valuation_2026-07.xlsx`` — July weighted-average-cost
    valuation: raw materials, finished goods, July purchases, recipe cost/kg.
  * ``stock_count_2026-07.xlsx`` / ``stock_count_2026-08.xlsx`` — physical
    counts by location (quantities only).
  * ``purchases_2026-08.pdf`` — purchases register for August.
  * ``suppliers_2026-09-02.pdf`` — supplier balances printed 02/09/2026.
  * ``returns_items_2026-07.pdf`` / ``returns_items_2026-08.pdf`` — sales
    returns with item, quantity and value.

Cleaning contract (every rule is counted and published in ``cleaning``):
  * Arabic presentation forms (PDF text) -> NFKC; bidi marks removed;
    Arabic-Indic digits -> ASCII.
  * Title / header / total / blank rows dropped (never data rows).
  * Item names normalised (ة/ه, ى/ي, أإآ/ا, spacing) only for MATCHING; the
    printed name is kept as ``name``.
  * Inventory / returns items are linked to the sales item codes of the
    dashboard's 04/09 snapshot by a (product, brand, size) signature. A link is "high" when the signature
    matches one sales item; otherwise the row stays unlinked — no guessing.
  * Every parsed total must equal the total printed in the source file
    (assert), so a parsing slip stops the build instead of reaching a chart.

Stdlib + openpyxl only (pdftotext from poppler for the PDFs):

    python3 -m src.ops_ingest          # from executive_dashboard/
"""
from __future__ import annotations

import json
import re
import subprocess
import unicodedata
from datetime import date
from pathlib import Path

APP_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = APP_DIR.parent
RAW = REPO_ROOT / "data" / "raw" / "ops"
SNAP_HTML = APP_DIR / "mobile" / "Abu_Hashem_Mobile_standalone.html"
OUT_JSON = REPO_ROOT / "data" / "processed" / "ops_data.json"

F_VAL_JUN = RAW / "inventory_valuation_2026-06.xlsx"
F_VAL_JUL = RAW / "inventory_valuation_2026-07.xlsx"
F_CNT = {"2026-07": RAW / "stock_count_2026-07.xlsx", "2026-08": RAW / "stock_count_2026-08.xlsx"}
F_PUR_AUG = RAW / "purchases_2026-08.pdf"
F_SUP = RAW / "suppliers_2026-09-02.pdf"
F_RET = {"2026-07": RAW / "returns_items_2026-07.pdf", "2026-08": RAW / "returns_items_2026-08.pdf"}

LOCATIONS = ["المصنع", "النفق", "النهضة", "المنوفية", "التغليف"]
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
BIDI = re.compile("[‎‏‪-‮⁦-⁩]")


class Log:
    """Counts what each cleaning rule touched (published, not hidden)."""

    def __init__(self):
        self.rules: dict[str, dict] = {}

    def hit(self, rule: str, title: str, example=None, n: int = 1):
        r = self.rules.setdefault(rule, {"id": rule, "title": title, "count": 0, "examples": []})
        r["count"] += n
        if example is not None and len(r["examples"]) < 5:
            r["examples"].append(example)

    def out(self) -> list:
        return list(self.rules.values())


LOG = Log()


# ---------------------------------------------------------------- text helpers
def clean_text(s) -> str:
    if s is None:
        return ""
    t = unicodedata.normalize("NFKC", str(s))
    t = BIDI.sub("", t).translate(DIGITS)
    return re.sub(r"\s+", " ", t).strip()


def norm_name(s) -> str:
    t = clean_text(s)
    for a, b in (("ة", "ه"), ("ى", "ي"), ("أ", "ا"), ("إ", "ا"), ("آ", "ا"), ("ـ", "")):
        t = t.replace(a, b)
    return re.sub(r"\s+", " ", t).strip()


def num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = clean_text(v).replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


TYPES = ["صويا", "سجق", "كفته", "برجر", "بانيه", "استربس", "شاورما", "بفتيك", "كبده", "مفروم", "ممبار",
         "بسطرمه", "لانشون", "هوت"]
BRANDS = [("الهنا", "HN"), ("اسبيشيال", "SP"), ("اسبشيال", "SP"), ("سبيشال", "SP"), ("اسباشيل", "SP"),
          ("ابوهاشم", "AH"), ("ابو هاشم", "AH"), ("الهاشم", "AH"), ("هنا", "HN"), ("SP", "SP")]
# a variant word on one side only means a different product -> no link
VARIANTS = ["حار", "بارد", "عرض", "مخصوص", "خام", "امريكي", "برازيلي", "دوفا"]
SIZE_WORDS = [("صغير", "S"), ("عائلي", "F")]


def signature(name: str):
    """(type, brand, size) used only to MATCH names; None when not comparable.

    * brand may be None (e.g. "بانيه 800 جرام") — it then only matches items
      that also carry no brand, never a branded one.
    * dough ("عجينة سجق") is its own type, so a raw paste is never linked to
      the finished product.
    * size: grams; a number < 10 is kilograms ("1.5 كيلو", "5ك"); the pack
      words صغير/عائلي are sizes too.
    """
    t = norm_name(name).replace("استبربس", "استربس")
    t = re.sub(r"(\d)\s*(جرام|جم|ج\b)", r"\1", t)
    ptype = next((p for p in TYPES if p in t), None)
    if "عجينه" in t:
        ptype = "عجينه" + ("-" + ptype if ptype else "")
    brand = next((code for word, code in BRANDS if word in t), None)
    size = next((code for word, code in SIZE_WORDS if word in t), None)
    if size is None:
        nums = re.findall(r"\d+(?:\.\d+)?", t)
        if nums:
            v = float(nums[0])
            size = int(round(v * 1000)) if v < 10 else int(v)
    var = tuple(v for v in VARIANTS if v in t)
    return (ptype, brand, size, var) if ptype and size else None


def load_items() -> list:
    """Sales item universe = the item codes in the dashboard's 04/09 snapshot.

    Read from the preserved original mobile file (the same lines every sales
    section of the dashboard analyses), so a link always points at a code the
    page can show sales for.
    """
    html = SNAP_HTML.read_text(encoding="utf-8")
    m = re.search(r'<script type="application/json" id="snap-dash-2026-09-04.json">(.*?)</script>', html, re.S)
    lines = json.loads(m.group(1).replace("<\\/", "</"))["lines"]
    names: dict[str, str] = {}
    for ln in lines:
        names.setdefault(str(ln["item_code"]), ln["item_name"])
    out = []
    for code, nm in names.items():
        sig = signature(nm)
        out.append({"code": code, "name": nm, "sigs": {sig} if sig else set()})
    return out


def sig_text(sig) -> str:
    return f"{sig[0]}/{sig[1] or '—'}/{sig[2]}" + (f" ({' '.join(sig[3])})" if sig[3] else "")


def link(name: str, items: list) -> dict:
    sig = signature(name)
    if not sig:
        return {"item_code": None, "match": "غير مربوط", "match_basis": "لا توقيع (نوع/علامة/حجم) قابل للمقارنة"}
    hits = [it for it in items if sig in it["sigs"]]
    if len(hits) == 1:
        return {"item_code": hits[0]["code"], "item_name_sales": hits[0]["name"], "match": "عالية",
                "match_basis": "توقيع مطابق " + sig_text(sig)}
    if len(hits) > 1:
        return {"item_code": None, "match": "غير مربوط",
                "match_basis": f"توقيع يطابق {len(hits)} أصناف — لا اختيار تخميني"}
    return {"item_code": None, "match": "غير مربوط",
            "match_basis": "لا صنف مبيعات بالتوقيع " + sig_text(sig)}


def is_junk_name(name: str) -> bool:
    """A name cell holding only a number ("0", "0.0") is a spreadsheet residue."""
    if name and num(name) is not None:
        LOG.hit("drop_numeric_name", "حذف خلايا اسم صنف تحوي رقمًا فقط (مثل «0») من ملفات الجرد", name)
        return True
    return False


def recipe_family(product: str, brand: str, items: list) -> list:
    """Sales codes of the recipe's (type, brand) family, every pack size.

    A recipe gives cost per kg of the mixture, not per pack, so it maps to the
    whole family; dough and variant items are never part of it."""
    t = norm_name(product + " " + brand)
    ptype = next((p for p in TYPES if p in t), None)
    br = next((code for word, code in BRANDS if word in t), None)
    out = []
    for it in items:
        for sig in it["sigs"]:
            if sig[0] == ptype and sig[1] == br and not sig[3]:
                out.append(it["code"])
    return sorted(set(out))


REGION_CANON = {"الجيزه": "الجيزة", "القاهره": "القاهرة", "القليوبيه": "القليوبية", "المنوفيه": "المنوفية",
                "السماعليه": "الإسماعيلية", "السماعيلية": "الإسماعيلية", "الاسماعيلية": "الإسماعيلية",
                "الشرقيه": "الشرقية", "الغربيه": "الغربية"}


# ---------------------------------------------------------------- PDF text
def pdf_lines(path: Path) -> list:
    raw = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, check=True).stdout
    return [clean_text(x) if not x.strip() else BIDI.sub("", unicodedata.normalize("NFKC", x)).translate(DIGITS)
            for x in raw.decode("utf-8").splitlines() if x.strip()]


def printed_total(lines: list, label_re: str) -> float:
    for ln in lines:
        if re.search(label_re, ln):
            m = re.findall(r"\d[\d,]*\.?\d*", ln)
            if m:
                return float(max(m, key=len).replace(",", ""))
    raise ValueError(f"printed total not found: {label_re}")


# ---------------------------------------------------------------- inventory
def parse_val_jun(items: list) -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(F_VAL_JUN, data_only=True).worksheets[0]
    raw, fg, total, sub = [], [], None, {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        a = clean_text(row[0])
        if a.startswith("إجمالى"):
            # printed totals sit in the value column of their block (I or M), never in B
            v = next((num(x) for x in reversed(row) if num(x) is not None), None)
            if a.startswith("إجمالى قيمة المخزون الكلى"):
                total = v
            elif "الخامات" in a:
                sub["raw"] = v
            elif "التام" in a:
                sub["fg"] = v
            LOG.hit("drop_header", "حذف صفوف العناوين والإجماليات والملاحظات من ملفات الجرد", a[:40])
            continue
        if not a or a == "اسم الصنف" or a.startswith("جرد") or a.startswith("ملاحظة"):
            if a:
                LOG.hit("drop_header", "حذف صفوف العناوين والإجماليات والملاحظات من ملفات الجرد", a[:40])
            elif any(v not in (None, "") for v in row[:9]):
                LOG.hit("drop_header", "حذف صفوف العناوين والإجماليات والملاحظات من ملفات الجرد", "صف بلا اسم")
        else:
            qtys = [num(row[i]) or 0.0 for i in range(1, 6)]
            raw.append({"name": a, "qty": dict(zip(LOCATIONS, qtys)), "qty_total": num(row[6]) or 0.0,
                        "price": num(row[7]), "value": num(row[8]) or 0.0})
        fgn = clean_text(row[9]) if len(row) > 9 else ""
        if fgn and fgn != "اسم الصنف" and not is_junk_name(fgn):
            rec = {"name": fgn, "qty": num(row[10]) or 0.0, "price": num(row[11]), "value": num(row[12]) or 0.0,
                   "note": clean_text(row[14]) or clean_text(row[13])}
            rec.update(link(fgn, items))
            fg.append(rec)
    raw_v = sum(r["value"] for r in raw)
    fg_v = sum(r["value"] for r in fg)
    assert total is not None and abs(raw_v + fg_v - total) < 1, (raw_v, fg_v, total)
    assert abs(raw_v - sub["raw"]) < 1 and abs(fg_v - sub["fg"]) < 1, (raw_v, fg_v, sub)
    return {"date": "2026-06-30", "raw": raw, "fg": fg, "raw_value": raw_v, "fg_value": fg_v, "total_printed": total}


def parse_val_jul(items: list) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(F_VAL_JUL, data_only=True)
    ws = wb["جرد الخامات آخر يوليو"]
    raw, raw_total = [], None
    for row in ws.iter_rows(min_row=5, values_only=True):
        a = clean_text(row[0])
        if a.startswith("إجمالى قيمة مخزون الخامات"):
            raw_total = next((num(x) for x in reversed(row) if num(x) is not None), None)
            continue
        if not a:
            continue
        raw.append({"name": a, "open_qty": num(row[1]) or 0.0, "open_price": num(row[2]), "open_value": num(row[3]) or 0.0,
                    "purch_qty": num(row[4]) or 0.0, "purch_value": num(row[5]) or 0.0, "wac": num(row[8]),
                    "qty": num(row[9]) or 0.0, "value": num(row[10]) or 0.0, "chg_vs_jun": num(row[11])})
    ws = wb["جرد المنتج التام آخر يوليو"]
    fg, fg_total, all_total = [], None, None
    for row in ws.iter_rows(min_row=4, values_only=True):
        a = clean_text(row[0])
        if a.startswith("إجمالى قيمة مخزون المنتج التام"):
            fg_total = next((num(x) for x in reversed(row) if num(x) is not None), None)
            continue
        if a.startswith("إجمالى قيمة المخزون الكلى"):
            all_total = next((num(x) for x in reversed(row) if num(x) is not None), None)
            continue
        if not a:
            continue
        rec = {"name": a, "qty": num(row[1]) or 0.0, "price_jun": num(row[2]), "price": num(row[3]),
               "chg": num(row[4]), "value": num(row[5]) or 0.0, "note": clean_text(row[6])}
        rec.update(link(a, items))
        fg.append(rec)
    ws = wb["مشتريات يوليو التفصيلية"]
    pur, pur_total = [], None
    for row in ws.iter_rows(min_row=4, values_only=True):
        a = clean_text(row[0])
        if any(clean_text(x).startswith("إجمالى مشتريات يوليو") for x in row[:4]):  # label sits in column D
            pur_total = next((num(x) for x in reversed(row) if num(x) is not None), None)
            continue
        if num(row[0]) is None:
            continue
        pur.append({"date": clean_text(row[1]).replace("/", "-"), "code": clean_text(row[2]), "name": clean_text(row[3]),
                    "qty": num(row[4]) or 0.0, "price": num(row[5]), "total": num(row[6]) or 0.0,
                    "stock_item": clean_text(row[7]), "confidence": clean_text(row[8])})
    ws = wb["سعر الكجم من الوصفات"]
    recipes = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        a = clean_text(row[0])
        if not a or num(row[4]) is None:
            continue
        rec = {"product": a, "brand": clean_text(row[1]), "batch_jun": num(row[2]), "batch_jul": num(row[3]),
               "kg_jun": num(row[4]), "kg_jul": num(row[5]), "chg": num(row[6])}
        recipes.append(rec)
    rv, fv, pv = (sum(r["value"] for r in raw), sum(r["value"] for r in fg), sum(r["total"] for r in pur))
    assert abs(rv - raw_total) < 1, (rv, raw_total)
    assert abs(fv - fg_total) < 1, (fv, fg_total)
    assert abs(rv + fv - all_total) < 1, (rv + fv, all_total)
    assert abs(pv - pur_total) < 1, (pv, pur_total)
    return {"date": "2026-07-31", "raw": raw, "fg": fg, "raw_value": rv, "fg_value": fv, "total_printed": all_total,
            "purchases": pur, "purchases_total": pv, "recipes": recipes}


def parse_count(path: Path, items: list) -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    raw, fg = [], []
    for row in ws.iter_rows(min_row=1, values_only=True):
        a = clean_text(row[0])
        if a and a != "اسم الصنف" and not a.startswith("جرد") and not is_junk_name(a):
            qtys = [num(row[i]) or 0.0 for i in range(1, 6)]
            raw.append({"name": a, "qty": dict(zip(LOCATIONS, qtys)), "qty_total": num(row[6]) or 0.0})
        fgn = clean_text(row[7]) if len(row) > 7 else ""
        if fgn and fgn != "اسم الصنف" and not is_junk_name(fgn):
            rec = {"name": fgn, "qty": num(row[8]) or 0.0}
            rec.update(link(fgn, items))
            fg.append(rec)
    return {"raw": raw, "fg": fg}


# ---------------------------------------------------------------- PDFs
ROW_PUR = re.compile(r"^\s*(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s+(\d{4}/\d{2}/\d{2})\s+([\d,]+(?:\.\d+)?)\s+"
                     r"([\d,]+(?:\.\d+)?)\s*(\d{6,})\s+(\d+)\s+(\S+)\s+(.+?)\s+(\d+)\s*$")


def parse_purchases_aug() -> dict:
    lines = pdf_lines(F_PUR_AUG)
    rows = []
    for ln in lines:
        m = ROW_PUR.match(ln)
        if not m:
            continue
        sale, price, d, total, qty, inv, code, unit, name, idx = m.groups()
        rows.append({"idx": int(idx), "date": d.replace("/", "-"), "invoice": inv, "code": code, "unit": unit,
                     "name": clean_text(name), "qty": float(qty.replace(",", "")), "price": float(price),
                     "total": float(total.replace(",", "")), "sale_price": float(sale)})
    printed = printed_total(lines, r"الجمالي\s*$|^\s*[\d,]+\.\d+\s+الجمالي")
    s = sum(r["total"] for r in rows)
    assert abs(s - printed) < 1, (s, printed, len(rows))
    assert [r["idx"] for r in rows] == list(range(1, len(rows) + 1)), "missing purchase rows"
    LOG.hit("pdf_text", "تحويل نص PDF (أشكال العرض العربية ← NFKC، إزالة علامات الاتجاه، الأرقام الهندية ← لاتينية)",
            "مشتريات أغسطس", n=len(rows))
    return {"month": "2026-08", "rows": rows, "total": s, "total_printed": printed}


ROW_SUP = re.compile(r"^\s*(?:(\d{8,})\s+)?([\d,]+\.\d+)\s+([\d,]+\.\d+)\s+(.+?)\s+(\d+)\s*$")


def parse_suppliers() -> dict:
    lines = pdf_lines(F_SUP)
    rows = []
    for ln in lines:
        m = ROW_SUP.match(ln)
        if not m:
            continue
        phone, debit, credit, name, idx = m.groups()
        rows.append({"idx": int(idx), "name": clean_text(name), "debit": float(debit.replace(",", "")),
                     "credit": float(credit.replace(",", "")), "has_phone": bool(phone)})
    tot_line = next(ln for ln in lines if "الجمالى" in ln and "دائن" in ln)
    printed = float(re.search(r"([\d,]+\.\d+)", tot_line).group(1).replace(",", ""))
    net = sum(r["credit"] - r["debit"] for r in rows)
    assert abs(net - printed) < 1, (net, printed, len(rows))
    assert [r["idx"] for r in rows] == list(range(1, len(rows) + 1)), "missing supplier rows"
    LOG.hit("pdf_text", "تحويل نص PDF (أشكال العرض العربية ← NFKC، إزالة علامات الاتجاه، الأرقام الهندية ← لاتينية)",
            "حسابات الموردين", n=len(rows))
    tiny = [r for r in rows if max(r["credit"], r["debit"]) < 50]
    if tiny:
        LOG.hit("supplier_dust", "أرصدة موردين أقل من 50 ج.م (فروق تقريب) تُعرض ولا تُحذف، وتُستبعد من باريتو",
                [r["name"] for r in tiny][:5], n=len(tiny))
    return {"as_of": "2026-09-02", "rows": rows, "net_credit": net, "total_printed": printed}


RET_HEAD = re.compile(r"^\s*(\d{4}/\d{1,2}/\d{1,2})\s+([\d,]+\.\d+)\s+([\d,]+\.\d+)\s+(.+?)\s+(\d{3,})\s+(\d+)\s*$")


def split_return_middle(mid: str):
    """unit, item, governorate, rep, customer.

    Columns are 3+ spaces apart in the pdftotext layout; a 2-space gap is the
    bidi wrap around digits *inside* a field ("كفته الهنا  800جرام",
    "السلمية  2م خليل"), so it never separates columns.
    """
    seg = [x for x in re.split(r"\s{3,}", mid.strip()) if x]
    if len(seg) != 5:
        return None
    return tuple(re.sub(r"\s+", " ", x) for x in seg)


def parse_returns(month: str, path: Path, items: list) -> dict:
    lines = pdf_lines(path)
    rows = []
    last = None
    for ln in lines:
        m = RET_HEAD.match(ln)
        if not m:
            t = ln.strip()
            # a wrapped item name continues on its own short line ("كيلو")
            if last is not None and t and len(t) <= 12 and not re.search(r"\d", t):
                last["item"] = clean_text(last["item"] + " " + t)
                LOG.hit("wrap_join", "دمج اسم صنف التف على سطر ثانٍ في PDF المرتجعات", last["item"])
            continue
        d, val, qty, mid, inv, idx = m.groups()
        parts = split_return_middle(mid)
        assert parts, ("unparsed return row", ln)
        unit, item, gov, rep, cust = parts
        y, mo, da = d.split("/")
        last = {"idx": int(idx), "date": f"{y}-{int(mo):02d}-{int(da):02d}", "value": float(val.replace(",", "")),
                "qty": float(qty.replace(",", "")), "unit": clean_text(unit), "item": clean_text(item),
                "region": clean_text(gov), "rep": clean_text(rep), "customer": clean_text(cust), "invoice": inv}
        rows.append(last)
    for r in rows:
        canon = REGION_CANON.get(r["region"])
        if canon:
            LOG.hit("region_canon", "توحيد كتابة أسماء المحافظات (الجيزه/الجيزة، السماعليه/الإسماعيلية …) للتجميع",
                    f'{r["region"]} ← {canon}')
            r["region_printed"], r["region"] = r["region"], canon
        if r["region"] == "...":
            LOG.hit("region_blank", "محافظة مطبوعة «...» في PDF المرتجعات تُسجَّل «غير محددة» (لا تخمين)", r["customer"])
            r["region"] = "غير محددة"
        r.update(link(r["item"], items))
    printed = printed_total(lines, r"اجمالي المرتجع")
    s = sum(r["value"] for r in rows)
    assert abs(s - printed) < 1, (month, s, printed, len(rows))
    assert sorted(r["idx"] for r in rows) == list(range(1, len(rows) + 1)), f"missing returns rows {month}"
    LOG.hit("pdf_text", "تحويل نص PDF (أشكال العرض العربية ← NFKC، إزالة علامات الاتجاه، الأرقام الهندية ← لاتينية)",
            f"مرتجعات {month}", n=len(rows))
    return {"month": month, "rows": rows, "total": s, "total_printed": printed}


# ---------------------------------------------------------------- build
def build() -> dict:
    items = load_items()
    jun = parse_val_jun(items)
    jul = parse_val_jul(items)
    for r in jul["recipes"]:
        r["sales_codes"] = recipe_family(r["product"], r["brand"], items)
        assert r["sales_codes"], ("recipe without a sales family", r["product"])
    counts = {m: parse_count(p, items) for m, p in F_CNT.items()}
    pur_aug = parse_purchases_aug()
    sup = parse_suppliers()
    ret = {m: parse_returns(m, p, items) for m, p in F_RET.items()}

    linked = [r for src in (jun["fg"], jul["fg"], counts["2026-07"]["fg"], counts["2026-08"]["fg"])
              for r in src] + [r for m in ret.values() for r in m["rows"]]
    hi = [r for r in linked if r["match"] == "عالية"]
    LOG.hit("link_items", "ربط أسماء المنتج التام والمرتجعات بأكواد أصناف المبيعات بتوقيع (نوع/علامة/حجم)",
            [f'{r.get("name") or r.get("item")} ← {r["item_code"]}' for r in hi[:5]], n=len(hi))
    unl = sorted({(r.get("name") or r.get("item")) for r in linked if r["match"] != "عالية"})
    LOG.hit("unlinked", "أسماء بقيت غير مربوطة (لا ربط تخميني)", unl[:5], n=len(unl))

    return {
        "meta": {
            "generated": date.today().isoformat(),
            "sources": [
                {"file": p.name, "kind": k} for p, k in [
                    (F_VAL_JUN, "جرد وتقييم آخر يونيو 2026"), (F_VAL_JUL, "تقييم مخزون آخر يوليو 2026 بمتوسط التكلفة"),
                    (F_CNT["2026-07"], "جرد كميات يوليو"), (F_CNT["2026-08"], "جرد كميات أغسطس"),
                    (F_PUR_AUG, "تقرير مشتريات أغسطس 2026"), (F_SUP, "حسابات الموردين 02/09/2026"),
                    (F_RET["2026-07"], "مرتجعات يوليو بالصنف"), (F_RET["2026-08"], "مرتجعات أغسطس بالصنف")]],
            "notes": [
                "تقييم المخزون بتكلفة الخامات والتعبئة فقط؛ لا يشمل أجورًا مباشرة ولا أعباء صناعية غير مباشرة "
                "(كما ينص ملف المصدر).",
                "جرد أغسطس كميات فقط بلا أسعار؛ قيمته غير متاحة.",
                "عمود «سعر البيع الحالي» في تقرير مشتريات أغسطس محفوظ كما طُبع (sale_price) ولا يدخل في أي حساب.",
            ],
        },
        "inventory": {"jun": jun, "jul": {k: v for k, v in jul.items() if k not in ("purchases", "recipes")},
                      "counts": counts, "recipes": jul["recipes"], "locations": LOCATIONS},
        "purchases": {"2026-07": {"rows": jul["purchases"], "total": jul["purchases_total"]}, "2026-08": pur_aug},
        "suppliers": sup,
        "returns_items": ret,
        "cleaning": LOG.out(),
        "unlinked_names": unl,
    }


def main() -> int:
    payload = build()
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    inv = payload["inventory"]
    print("inventory jun", round(inv["jun"]["raw_value"] + inv["jun"]["fg_value"], 2),
          "jul", round(inv["jul"]["raw_value"] + inv["jul"]["fg_value"], 2))
    print("purchases", {m: round(v["total"], 2) for m, v in payload["purchases"].items()},
          "suppliers net credit", round(payload["suppliers"]["net_credit"], 2))
    print("returns", {m: (len(v["rows"]), round(v["total"], 2)) for m, v in payload["returns_items"].items()})
    for r in payload["cleaning"]:
        print(" -", r["id"], r["count"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
