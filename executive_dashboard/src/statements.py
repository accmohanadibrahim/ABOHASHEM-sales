"""
Financial statements (balance sheet + monthly P&L) for the finance screens.

Sources (copied unchanged into ``data/raw/``):
  * ``financial_statements_pop.xlsx``  — B.S. / I.S. / other-current-assets /
    CASH FLOW / RATIOS workbook. Balance-sheet columns 31/7/2025, 31/8/2025,
    31/12/2025 and an incomplete 28/2/2026.
  * ``income_statement_detail_2025H2.xlsx`` — detailed monthly income statement
    Jul → Dec 2025 (one sheet per month, expense lines itemised).

Rules (same anti-fabrication contract as the rest of the dashboard):
  * Read rows by their account CODE / label, never by fixed position.
  * A balance-sheet column is used only if assets == liabilities + equity
    (±1 EGP). Unbalanced columns are excluded and reported as a quality flag.
  * The workbook's own CASH FLOW / OPS / RATIOS sheets are NOT consumed: they
    are derived from the unbalanced Feb-2026 column. Cash flow and ratios are
    recomputed here from balanced balance sheets + the monthly P&L, and every
    formula is published in ``kpi_dictionary``.

Stdlib + openpyxl only, so it runs (and is tested) without Polars:

    python3 -m src.statements          # from executive_dashboard/
"""
from __future__ import annotations

import csv
import json
import re
from datetime import date
from pathlib import Path

APP_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = APP_DIR.parent
RAW = REPO_ROOT / "data" / "raw"
F_POP = RAW / "financial_statements_pop.xlsx"
F_IS_DETAIL = RAW / "income_statement_detail_2025H2.xlsx"
F_MONTHLY_INVOICES = REPO_ROOT / "data" / "processed" / "eda_monthly_series.csv"
F_BUDGET = REPO_ROOT / "data" / "budget" / "budget_2026.csv"
OUT_JSON = REPO_ROOT / "data" / "processed" / "financial_statements.json"
OUT_JS = APP_DIR / "statements_data.js"

BALANCE_TOL = 1.0            # EGP — assets vs liabilities+equity
RECON_TOL_PCT = 0.02         # invoice sales vs statement net sales
ITEM_SUM_TOL = 1.0           # itemised expenses vs section total

# Alert thresholds (single reviewable place, like src/config.py)
ALERTS = {
    "gross_margin_drop_pts": 2.0,      # month-on-month gross-margin drop
    "admin_pct_of_sales_max": 0.15,    # G&A ÷ net sales
    "drawings_pct_of_profit_max": 0.50,
    "dio_days_max": 60,
    "sales_recon_pct_max": RECON_TOL_PCT,
}

# --- Balance sheet: workbook code -> key -----------------------------------
BS_CODES = {
    "01": "cash", "02": "time_deposits", "03": "receivables", "04": "inventory",
    "05": "prepaid", "06": "other_current_assets", "07": "due_from_affiliates",
    "08": "total_current_assets", "09": "net_fixed_assets",
    "10": "investments_subsidiaries", "11": "lt_investments", "12": "tax_receivable",
    "13": "other_noncurrent_assets", "14": "intangibles", "15": "total_noncurrent_assets",
    "16": "total_assets", "17": "st_bank_debt", "18": "payables",
    "19": "accrued_expenses", "20": "other_credit_balances", "21": "dividends_payable",
    "22": "provisions_current", "23": "due_to_affiliates", "24": "current_ltd",
    "25": "total_current_liabilities", "26": "initiative_loan", "27": "car_loan",
    "28": "shareholder_loans", "30": "tax_provision", "31": "other_provisions",
    "32": "total_lt_liabilities", "33": "total_liabilities",
    "35": "paid_in_capital", "36": "partners_current", "37": "reserves",
    "38": "treasury_shares", "39": "retained_earnings", "40": "net_profit_bs",
    "41": "total_equity", "42": "total_liabilities_equity",
}
BS_LABELS_AR = {
    "cash": "النقدية", "time_deposits": "ودائع لأجل", "receivables": "العملاء وأوراق القبض",
    "inventory": "المخزون", "prepaid": "مصروفات مقدمة", "other_current_assets": "أصول متداولة أخرى",
    "total_current_assets": "إجمالي الأصول المتداولة", "net_fixed_assets": "صافي الأصول الثابتة",
    "investments_subsidiaries": "استثمارات في شركات تابعة",
    "total_noncurrent_assets": "إجمالي الأصول غير المتداولة", "total_assets": "إجمالي الأصول",
    "payables": "موردون وأوراق دفع", "accrued_expenses": "مصروفات مستحقة",
    "total_current_liabilities": "إجمالي الخصوم المتداولة", "initiative_loan": "قرض المبادرة",
    "car_loan": "قرض السيارة", "shareholder_loans": "قروض الشركاء (أ. خالد والحاج أسامة)",
    "total_lt_liabilities": "إجمالي الالتزامات طويلة الأجل", "total_liabilities": "إجمالي الالتزامات",
    "paid_in_capital": "رأس المال المدفوع", "partners_current": "جاري الشركاء",
    "retained_earnings": "أرباح مرحلة", "net_profit_bs": "صافي الربح (كما بالميزانية)",
    "total_equity": "إجمالي حقوق الملكية", "total_liabilities_equity": "إجمالي الالتزامات وحقوق الملكية",
}
INTEREST_DEBT = ("st_bank_debt", "current_ltd", "initiative_loan", "car_loan", "shareholder_loans")


def _num(v) -> float:
    return float(v) if isinstance(v, (int, float)) else 0.0


def _date_key(txt) -> str | None:
    """'31/12/2025' or datetime -> '2025-12-31'."""
    if hasattr(txt, "year"):
        return f"{txt.year:04d}-{txt.month:02d}-{txt.day:02d}"
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", str(txt or ""))
    return f"{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}" if m else None


# --------------------------------------------------------------------------- #
#  Balance sheet                                                              #
# --------------------------------------------------------------------------- #
def parse_balance_sheets(wb) -> tuple[dict, list]:
    ws = wb["B.S."]
    # date header row: column A == "DATE (Day, Month, Year)"
    date_cols = {}
    for row in ws.iter_rows():
        if str(row[0].value or "").strip().startswith("DATE"):
            for c in row[3:]:
                k = _date_key(c.value)
                if k:
                    date_cols[c.column] = k
            break
    sheets = {d: {} for d in date_cols.values()}
    for row in ws.iter_rows():
        code = str(row[0].value or "").strip()
        if code.isdigit() and len(code) == 1:
            code = "0" + code
        key = BS_CODES.get(code)
        if not key:
            continue
        for col, d in date_cols.items():
            sheets[d][key] = round(_num(ws.cell(row=row[0].row, column=col).value), 2)

    balanced, flags = {}, []
    for d, bs in sorted(sheets.items()):
        diff = bs.get("total_assets", 0) - bs.get("total_liabilities_equity", 0)
        bs["balance_diff"] = round(diff, 2)
        if abs(diff) <= BALANCE_TOL and bs.get("total_assets"):
            balanced[d] = bs
        else:
            flags.append({
                "id": f"bs_unbalanced_{d}", "severity": "high",
                "title": f"عمود الميزانية {d} غير متوازن — مُستبعد",
                "detail": (f"إجمالي الأصول {bs.get('total_assets', 0):,.0f} مقابل الالتزامات وحقوق الملكية "
                           f"{bs.get('total_liabilities_equity', 0):,.0f} (فرق {diff:,.0f}). العمود ناقص "
                           "(لا نقدية ولا عملاء ولا مخزون)، ولذلك لا يُعتمد شيت CASH FLOW/OPS/RATIOS "
                           "في الملف لأنه محسوب منه؛ التدفق النقدي والنسب يُعاد حسابها هنا."),
                "action": "استكمال عمود الميزانية أو حذفه من الملف",
            })
    for bs in balanced.values():
        bs["interest_bearing_debt"] = round(sum(bs.get(k, 0) for k in INTEREST_DEBT), 2)
        bs["cash_and_deposits"] = round(bs.get("cash", 0) + bs.get("time_deposits", 0), 2)
        bs["net_debt"] = round(bs["interest_bearing_debt"] - bs["cash_and_deposits"], 2)
        bs["working_capital"] = round(bs["total_current_assets"] - bs["total_current_liabilities"], 2)
        cl = bs["total_current_liabilities"] or None
        bs["current_ratio"] = round(bs["total_current_assets"] / cl, 2) if cl else None
        bs["quick_ratio"] = round((bs["total_current_assets"] - bs["inventory"]) / cl, 2) if cl else None
        bs["cash_ratio"] = round(bs["cash_and_deposits"] / cl, 2) if cl else None
        eq = bs["total_equity"] or None
        bs["liabilities_to_equity"] = round(bs["total_liabilities"] / eq, 3) if eq else None
        bs["debt_to_equity"] = round(bs["interest_bearing_debt"] / eq, 3) if eq else None
        bs["equity_ratio"] = round(bs["total_equity"] / bs["total_assets"], 3)
    return balanced, flags


# --------------------------------------------------------------------------- #
#  Other current assets detail (sheet 2) — custody / advances / wallets        #
# --------------------------------------------------------------------------- #
OCA_SUBHEADERS = {"سلفيات طويلة الاجل", "عهد ثابتة", "عهد فودافون كاش وانستا",
                  "عهد سفريات وصيانة", "قروض", "قروض قصيرة", "أصول متداولة أخرى 3"}
OCA_BLOCK_DATES = {  # header text in row 2/3 -> period
    "ارصدة نهاية 7": "2025-07-31", "ارصدة نهاية 8": "2025-08-31",
    "ارصدة نهاية 9": "2025-09-30", "ارصدة نهاية 12 2025.": "2025-12-31",
    "ارصدة نهاية فبراير  2026": "2026-02-28",
}


def parse_other_current_assets(wb) -> dict:
    ws = wb.worksheets[1]
    starts = {}
    for r in (2, 3):
        for c in ws[r]:
            d = OCA_BLOCK_DATES.get(str(c.value or "").strip()) or OCA_BLOCK_DATES.get(str(c.value or ""))
            if d and d not in starts.values():
                starts[c.column] = d
    out = {}
    for start, d in starts.items():
        items = []
        for lab_col in (start, start + 3):
            cat = "أصول متداولة أخرى"
            for r in range(3, 45):
                lab = ws.cell(row=r, column=lab_col).value
                val = ws.cell(row=r, column=lab_col + 1).value
                if not isinstance(lab, str) or not lab.strip():
                    continue
                name = lab.strip()
                if name in OCA_SUBHEADERS and not isinstance(val, (int, float)):
                    cat = name
                    continue
                if isinstance(val, (int, float)):
                    items.append({"category": cat, "name": name, "amount": round(float(val), 2)})
        out[d] = items
    return out


# --------------------------------------------------------------------------- #
#  Monthly income statements                                                  #
# --------------------------------------------------------------------------- #
IS_SECTIONS = [  # (key, regex on the section label)
    ("net_sales", r"صافي المبيعات"), ("cogs", r"تكلفة المبيعات"), ("gross_profit", r"مجمل الربح"),
    ("admin", r"مصروفات إدارية"), ("selling", r"مصروفات بيعية"), ("finance", r"مصاريف تمويلية"),
    ("net_profit_pre_dep", r"صافى الربح|صافي الربح"),
]
MONTH_FROM_TITLE = re.compile(r"(\d{1,2}|[٠-٩]{1,2})\s*سنة")
AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


def _sheet_month(title: str) -> str | None:
    m = MONTH_FROM_TITLE.search(title.replace("سنه", "سنة"))
    if not m:
        return None
    return f"2025-{int(m[1].translate(AR_DIGITS)):02d}"


def _parse_is_sheet(ws) -> dict:
    """Section labels sit in column L, item names in L+1, item values in L+2 and
    section totals in L+3 (sales/COGS/gross values in L+2). A total may sit on
    the next row. Itemised lines are kept for the expense breakdown."""
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    label_col = next(col for (r, col), v in cells.items()
                     if isinstance(v, str) and "صافي المبيعات" in v)

    def first_num(r, cols):
        for col in cols:
            v = cells.get((r, col))
            if isinstance(v, (int, float)):
                return float(v)
        return None

    res, items, section, sec_rows = {}, {}, None, {}
    for r in sorted({r for r, _ in cells}):
        lab = cells.get((r, label_col)) or cells.get((r, label_col + 1))
        lab = lab if isinstance(lab, str) else ""
        hit = next((k for k, rx in IS_SECTIONS if re.search(rx, lab)), None)
        if hit:
            section, sec_rows[hit] = hit, r
            right = list(range(label_col + 1, label_col + 6))
            if hit in ("net_sales", "cogs", "gross_profit"):
                v = first_num(r, [label_col + 2, label_col + 3])
            else:
                # total in the totals column on the label row, else the row
                # below, else (net profit only) the row above the label.
                v = first_num(r, [label_col + 3])
                if v is None:
                    v = first_num(r + 1, [label_col + 3])
                if v is None and hit == "net_profit_pre_dep":
                    v = first_num(r - 1, [label_col + 3]) or first_num(r, right)
            res[hit] = v
            # An item printed on a section's header row continues the PREVIOUS
            # section (e.g. Sept salaries sit on the selling header but belong
            # to G&A — the G&A total only reconciles that way).
            prev = list(sec_rows)[-2] if len(sec_rows) > 1 else None
            nm = cells.get((r, label_col + 1))
            iv = cells.get((r, label_col + 2))
            if prev in ("admin", "selling", "finance") and isinstance(nm, str) and nm.strip() \
                    and isinstance(iv, (int, float)):
                items.setdefault(prev, []).append({"name": nm.strip(), "amount": float(iv)})
            continue
        if section in ("admin", "selling", "finance"):
            nm = cells.get((r, label_col + 1))
            iv = cells.get((r, label_col + 2))
            if isinstance(nm, str) and nm.strip() and isinstance(iv, (int, float)):
                items.setdefault(section, []).append({"name": nm.strip(), "amount": float(iv)})
    return {"totals": res, "items": items}


def parse_income_detail(wb) -> tuple[dict, list]:
    months, flags, item_gaps = {}, [], []
    for ws in wb.worksheets:
        m = _sheet_month(ws.title)
        if not m:
            continue
        p = _parse_is_sheet(ws)
        t = {k: (v if v is not None else None) for k, v in p["totals"].items()}
        for k in ("admin", "selling", "finance"):
            if t.get(k) is None:
                t[k] = round(sum(i["amount"] for i in p["items"].get(k, [])), 2)
        gp_calc = (t["net_sales"] or 0) - (t["cogs"] or 0)
        np_calc = gp_calc - t["admin"] - t["selling"] - (t.get("finance") or 0)
        other = 0.0
        if t.get("net_profit_pre_dep") is None:
            t["net_profit_pre_dep"] = round(np_calc, 2)
            flags.append({"id": f"is_np_computed_{m}", "severity": "low",
                          "title": f"صافي ربح {m} غير مكتوب في الورقة — محسوب",
                          "detail": "مجمل الربح − الإدارية − البيعية − التمويلية.",
                          "action": "كتابة الرقم في قائمة الدخل التفصيلية"})
        else:
            other = round(np_calc - t["net_profit_pre_dep"], 2)   # unexplained gap (e.g. FX)
        t["other_items"] = other
        for k in ("admin", "selling"):
            s = round(sum(i["amount"] for i in p["items"].get(k, [])), 2)
            if p["items"].get(k) and abs(s - t[k]) > ITEM_SUM_TOL:
                item_gaps.append(f"{m} {'الإدارية' if k == 'admin' else 'البيعية'}: البنود {s:,.0f} / الإجمالي {t[k]:,.0f}")
        months[m] = {"source": "income_statement_detail_2025H2.xlsx", **t,
                     "expense_items": p["items"]}
    if item_gaps:
        flags.append({"id": "is_item_totals", "severity": "low",
                      "title": "مجموع بنود المصروفات لا يساوي إجمالي القسم في بعض الشهور",
                      "detail": "؛ ".join(item_gaps) + ". الداشبورد يعتمد الإجمالي المكتوب في القائمة.",
                      "action": "مراجعة معادلة إجمالي القسم في قائمة الدخل التفصيلية"})
    return dict(sorted(months.items())), flags


def parse_pop_is(wb) -> dict:
    """I.S. sheet of the pop workbook: one column per month (codes 43..61)."""
    ws = wb["I.S."]
    cols = {}
    for row in ws.iter_rows():
        for c in row[3:]:
            k = _date_key(c.value)
            if k:
                cols[c.column] = k[:7]
        if cols:
            break
    codes = {"43": "net_sales", "44": "cogs", "45": "admin", "46": "ebitda_as_labelled",
             "47": "depreciation", "52": "other_expenses", "54": "fx_loss",
             "57": "finance", "61": "net_profit"}
    out = {m: {} for m in cols.values()}
    seen = set()
    for row in ws.iter_rows():
        key = codes.get(str(row[0].value or "").strip())
        if key and key not in seen:     # code 61 recurs in the net-worth reconciliation block
            seen.add(key)
            for col, m in cols.items():
                out[m][key] = round(_num(ws.cell(row=row[0].row, column=col).value), 2)
    return out


# --------------------------------------------------------------------------- #
#  Derived: P&L ratios, cash flow, working capital, reconciliation, flags     #
# --------------------------------------------------------------------------- #
def enrich_income(months: dict, pop_is: dict) -> dict:
    annual_dep = pop_is.get("2025-12", {}).get("depreciation", 0.0)
    for m, t in months.items():
        ns = t["net_sales"] or 0
        t["gross_profit"] = round(ns - (t["cogs"] or 0), 2)
        t["gross_margin"] = round(t["gross_profit"] / ns, 4) if ns else None
        t["ebitda"] = round(t["gross_profit"] - t["admin"] - t["selling"], 2)
        t["ebitda_margin"] = round(t["ebitda"] / ns, 4) if ns else None
        t["admin_pct"] = round(t["admin"] / ns, 4) if ns else None
        t["selling_pct"] = round(t["selling"] / ns, 4) if ns else None
        t["depreciation_booked"] = pop_is.get(m, {}).get("depreciation", 0.0)
        t["net_profit_after_dep"] = round(t["net_profit_pre_dep"] - t["depreciation_booked"], 2)
        # Annual depreciation booked once in December — spread evenly for a
        # comparable monthly view (labelled as an allocation estimate).
        t["depreciation_allocated_est"] = round(annual_dep / 12, 2) if annual_dep else None
        t["net_profit_dep_allocated_est"] = (round(t["net_profit_pre_dep"] - annual_dep / 12, 2)
                                             if annual_dep else None)
    return months


def cash_flows(bs: dict, months: dict) -> list:
    """Indirect cash flow between consecutive balanced balance sheets."""
    dates = sorted(bs)
    out = []
    for a, b in zip(dates, dates[1:]):
        pa, pb = bs[a], bs[b]
        period = [m for m in months if a[:7] < m <= b[:7]]
        ni = sum(months[m]["net_profit_after_dep"] for m in period)
        dep = sum(months[m]["depreciation_booked"] for m in period)
        d = lambda k: pb.get(k, 0) - pa.get(k, 0)  # noqa: E731
        wc_assets = -(d("receivables") + d("inventory") + d("prepaid") + d("other_current_assets"))
        wc_liab = d("payables") + d("accrued_expenses") + d("other_credit_balances")
        cfo = ni + dep + wc_assets + wc_liab
        capex = d("net_fixed_assets") + dep
        cfi = -capex - d("investments_subsidiaries")
        debt = sum(d(k) for k in INTEREST_DEBT)
        equity_other = d("total_equity") - ni        # drawings / injections (plug, labelled)
        cff = debt + equity_other
        dcash = d("cash_and_deposits")
        out.append({
            "from": a, "to": b, "months": period,
            "net_income": round(ni, 2), "depreciation": round(dep, 2),
            "change_receivables": round(-d("receivables"), 2), "change_inventory": round(-d("inventory"), 2),
            "change_other_current_assets": round(-(d("other_current_assets") + d("prepaid")), 2),
            "change_payables": round(d("payables"), 2),
            "change_accrued": round(d("accrued_expenses") + d("other_credit_balances"), 2),
            "cfo": round(cfo, 2), "capex": round(capex, 2), "cfi": round(cfi, 2),
            "debt_change": round(debt, 2), "equity_other_movements": round(equity_other, 2),
            "cff": round(cff, 2), "net_change": round(cfo + cfi + cff, 2),
            "actual_change_cash_deposits": round(dcash, 2),
            "check_diff": round(cfo + cfi + cff - dcash, 2),
        })
    return out


def working_capital(bs: dict, months: dict, as_of: str) -> dict | None:
    """DSO / DIO / DPO on the last 3 months before the balance-sheet date."""
    b = bs.get(as_of)
    win = [m for m in months if m <= as_of[:7]][-3:]
    if not b or len(win) < 3:
        return None
    days = sum({"01": 31, "02": 28, "03": 31, "04": 30, "05": 31, "06": 30, "07": 31, "08": 31,
                "09": 30, "10": 31, "11": 30, "12": 31}[m[5:]] for m in win)
    sales = sum(months[m]["net_sales"] for m in win)
    cogs = sum(months[m]["cogs"] for m in win)
    dso = b["receivables"] / sales * days
    dio = b["inventory"] / cogs * days
    dpo = b["payables"] / cogs * days
    return {"as_of": as_of, "window_months": win, "window_days": days,
            "sales_window": round(sales, 2), "cogs_window": round(cogs, 2),
            "dso": round(dso, 1), "dio": round(dio, 1), "dpo": round(dpo, 1),
            "ccc": round(dso + dio - dpo, 1),
            "cash_tied_receivables": b["receivables"], "cash_tied_inventory": b["inventory"],
            "funded_by_payables": b["payables"]}


def sales_reconciliation(months: dict) -> list:
    inv = {}
    if F_MONTHLY_INVOICES.exists():
        with open(F_MONTHLY_INVOICES, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                inv[r["month"]] = float(r["revenue"])
    rows = []
    for m, t in months.items():
        iv = inv.get(m)
        diff = None if iv is None else t["net_sales"] - iv
        rows.append({"month": m, "statement_net_sales": t["net_sales"], "invoice_sales": iv,
                     "diff": round(diff, 2) if diff is not None else None,
                     "diff_pct": round(diff / iv, 4) if iv else None,
                     "ok": (abs(diff / iv) <= RECON_TOL_PCT) if iv else None})
    return rows


def load_budget() -> dict:
    out = {}
    if F_BUDGET.exists():
        with open(F_BUDGET, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if (r.get("amount") or "").strip():
                    out.setdefault(r["month"], {})[r["line"]] = float(r["amount"])
    return out


def quality_flags(bs, months, pop_is, oca, recon) -> list:
    flags = []
    dec_pop = pop_is.get("2025-12", {})
    dec = months.get("2025-12", {})
    if dec_pop.get("depreciation"):
        flags.append({
            "id": "dep_booked_december", "severity": "medium",
            "title": "إهلاك السنة كاملًا مُحمَّل على ديسمبر",
            "detail": (f"ورقة I.S. تحمّل ديسمبر إهلاكًا قدره {dec_pop['depreciation']:,.0f} فتنقلب النتيجة "
                       f"إلى {dec_pop.get('net_profit', 0):,.0f}، بينما صافي ربح ديسمبر التشغيلي قبل الإهلاك "
                       f"{dec.get('net_profit_pre_dep', 0):,.0f}. يعرض الداشبورد الرقمين، ويوزّع الإهلاك "
                       "شهريًا كتقدير موسوم للمقارنة."),
            "action": "تأكيد أن الإهلاك سنوي، وتحميله شهريًا من 2026",
        })
    b12 = bs.get("2025-12-31")
    if b12:
        h2 = sum(t["net_profit_after_dep"] for m, t in months.items() if m >= "2025-07")
        flags.append({
            "id": "bs_net_profit_mismatch", "severity": "medium",
            "title": "«صافي الربح» في الميزانية لا يطابق قوائم الدخل الشهرية",
            "detail": (f"الميزانية 31/12/2025: {b12['net_profit_bs']:,.0f}. مجموع صافي الربح بعد الإهلاك "
                       f"يوليو–ديسمبر من القوائم الشهرية: {h2:,.0f}. الفترة التي يغطيها رقم الميزانية غير مذكورة."),
            "action": "توضيح الفترة التي يغطيها صافي الربح في الميزانية",
        })
    jul_pop, jul = pop_is.get("2025-07", {}), months.get("2025-07", {})
    if jul_pop and jul and abs(jul_pop.get("admin", 0) - jul.get("admin", 0)) > 1:
        flags.append({
            "id": "jul_admin_mismatch", "severity": "low",
            "title": "مصروفات يوليو 2025 الإدارية مختلفة بين الملفين",
            "detail": (f"ملف القوائم المالية {jul_pop['admin']:,.0f}، القائمة التفصيلية {jul['admin']:,.0f} "
                       f"(فرق {jul['admin'] - jul_pop['admin']:,.0f}). الداشبورد يعتمد القائمة التفصيلية."),
            "action": "اعتماد رقم واحد",
        })
    if dec_pop.get("ebitda_as_labelled"):
        flags.append({
            "id": "operating_profit_label", "severity": "low",
            "title": "«أرباح التشغيل» في ملف القوائم محسوبة قبل الإهلاك",
            "detail": "المبيعات − التكلفة − الإدارية فقط؛ تُعرض في الداشبورد باسم EBITDA بعد طرح البيعية أيضًا.",
            "action": "تسمية البند «ربح التشغيل قبل الإهلاك»",
        })
    dec_items = oca.get("2025-12-31", [])
    prov = [i for i in dec_items if "مخصص ديون" in i["name"] and i["amount"]]
    if prov and b12:
        p = prov[0]["amount"]
        flags.append({
            "id": "bad_debt_provision_placement", "severity": "medium",
            "title": "مخصص الديون المعدومة مُدرج ضمن الأصول المتداولة الأخرى",
            "detail": (f"المخصص {p:,.0f} يجب أن يُطرح من العملاء. صافي العملاء بعد المخصص "
                       f"{b12['receivables'] - p:,.0f} مقابل {b12['receivables']:,.0f} في الميزانية."),
            "action": "عرض المخصص مطروحًا من رصيد العملاء",
        })
    if b12 and dec_items:
        s = sum(i["amount"] for i in dec_items)
        flags.append({
            "id": "oca_detail_vs_bs", "severity": "medium",
            "title": "تفصيل الأصول المتداولة الأخرى لا يطابق رصيد الميزانية",
            "detail": (f"مجموع بنود ورقة التفصيل في 12/2025: {s:,.0f}، والرصيد في الميزانية "
                       f"{b12['other_current_assets']:,.0f} (فرق {s - b12['other_current_assets']:,.0f}). "
                       "الورقة تخلط سُلفًا وعُهدًا ومحافظ إلكترونية ورصيد قرض ومخصصًا."),
            "action": "تسوية البند وفصل كل نوع في حساب مستقل",
        })
    neg = [i for items in oca.values() for i in items if i["amount"] < 0]
    if neg:
        flags.append({
            "id": "negative_custody", "severity": "medium",
            "title": "أرصدة عُهد سالبة",
            "detail": "، ".join(f"{i['name']} {i['amount']:,.0f}" for i in neg[:6]),
            "action": "تسوية العهدة مع صاحبها",
        })
    for r in recon:
        if r["ok"] is False:
            flags.append({
                "id": f"sales_recon_{r['month']}", "severity": "medium",
                "title": f"مبيعات {r['month']}: القوائم لا تطابق الفواتير",
                "detail": (f"القوائم {r['statement_net_sales']:,.0f}، الفواتير {r['invoice_sales']:,.0f} "
                           f"(فرق {r['diff_pct'] * 100:.1f}%)."),
                "action": "تفسير الفرق (مرتجعات/خصومات/فواتير لاحقة)",
            })
    return flags


def alerts(bs, months, wc, cfs) -> list:
    out = []
    ms = list(months)
    for a, b in zip(ms, ms[1:]):
        ga, gb = months[a]["gross_margin"], months[b]["gross_margin"]
        if ga is not None and gb is not None and (ga - gb) * 100 > ALERTS["gross_margin_drop_pts"]:
            out.append({"month": b, "kind": "margin", "severity": "high",
                        "problem": f"هامش مجمل الربح هبط من {ga * 100:.1f}% إلى {gb * 100:.1f}%",
                        "value": round((gb - ga) * 100, 1), "owner": "مدير المصنع / المشتريات",
                        "action": "مراجعة أسعار الخامات وتكلفة الكيلو وتسعير البيع لهذا الشهر"})
    for m, t in months.items():
        if t["admin_pct"] and t["admin_pct"] > ALERTS["admin_pct_of_sales_max"]:
            out.append({"month": m, "kind": "opex", "severity": "medium",
                        "problem": f"المصروفات الإدارية {t['admin_pct'] * 100:.1f}% من المبيعات",
                        "value": t["admin"], "owner": "المدير المالي",
                        "action": "مراجعة أكبر بنود المصروفات الإدارية في الشهر"})
    if wc and wc["dio"] > ALERTS["dio_days_max"]:
        out.append({"month": wc["as_of"][:7], "kind": "inventory", "severity": "medium",
                    "problem": f"أيام المخزون {wc['dio']:.0f} يومًا", "value": wc["dio"],
                    "owner": "مدير المخازن", "action": "تحديد الأصناف الراكدة"})
    for cf in cfs:
        if cf["net_income"] > 0 and -cf["equity_other_movements"] > ALERTS["drawings_pct_of_profit_max"] * cf["net_income"]:
            out.append({"month": cf["to"][:7], "kind": "drawings", "severity": "high",
                        "problem": (f"مسحوبات/حركة جاري الشركاء {-cf['equity_other_movements']:,.0f} "
                                    f"مقابل صافي ربح الفترة {cf['net_income']:,.0f}"),
                        "value": -cf["equity_other_movements"], "owner": "الشركاء",
                        "action": "وضع سقف للمسحوبات مرتبط بالتدفق النقدي التشغيلي"})
    return out


KPI_DICTIONARY = [
    {"kpi": "هامش مجمل الربح", "formula": "(صافي المبيعات − تكلفة المبيعات) ÷ صافي المبيعات",
     "source": "قائمة الدخل التفصيلية 7–12/2025", "period": "شهري"},
    {"kpi": "EBITDA", "formula": "مجمل الربح − المصروفات الإدارية − المصروفات البيعية",
     "source": "قائمة الدخل التفصيلية", "period": "شهري"},
    {"kpi": "صافي الربح قبل الإهلاك", "formula": "EBITDA − مصاريف تمويلية − بنود أخرى (فروق عملة)",
     "source": "قائمة الدخل التفصيلية", "period": "شهري"},
    {"kpi": "صافي الربح بعد الإهلاك", "formula": "صافي الربح قبل الإهلاك − الإهلاك المحمَّل",
     "source": "ملف القوائم المالية، ورقة I.S. كود 47", "period": "شهري"},
    {"kpi": "رأس المال العامل", "formula": "الأصول المتداولة − الخصوم المتداولة",
     "source": "ورقة B.S. كود 08 − كود 25", "period": "في تاريخ الميزانية"},
    {"kpi": "صافي الدين", "formula": "القروض بفوائد (المبادرة + السيارة + الشركاء + البنوك) − (النقدية + الودائع)",
     "source": "ورقة B.S. أكواد 17/24/26/27/28 − 01/02", "period": "في تاريخ الميزانية"},
    {"kpi": "نسبة التداول", "formula": "الأصول المتداولة ÷ الخصوم المتداولة", "source": "B.S.",
     "period": "في تاريخ الميزانية"},
    {"kpi": "النسبة السريعة", "formula": "(الأصول المتداولة − المخزون) ÷ الخصوم المتداولة", "source": "B.S.",
     "period": "في تاريخ الميزانية"},
    {"kpi": "أيام التحصيل DSO", "formula": "رصيد العملاء ÷ مبيعات آخر 3 أشهر × عدد أيامها",
     "source": "B.S. كود 03 + قائمة الدخل", "period": "آخر 3 أشهر"},
    {"kpi": "أيام المخزون DIO", "formula": "المخزون ÷ تكلفة مبيعات آخر 3 أشهر × عدد أيامها",
     "source": "B.S. كود 04 + قائمة الدخل", "period": "آخر 3 أشهر"},
    {"kpi": "أيام السداد DPO", "formula": "الموردون ÷ تكلفة مبيعات آخر 3 أشهر × عدد أيامها",
     "source": "B.S. كود 18 + قائمة الدخل", "period": "آخر 3 أشهر"},
    {"kpi": "دورة التحويل النقدي", "formula": "DSO + DIO − DPO", "source": "محسوب", "period": "آخر 3 أشهر"},
    {"kpi": "التدفق النقدي التشغيلي",
     "formula": "صافي الربح + الإهلاك ± التغير في العملاء والمخزون والأصول المتداولة الأخرى والموردين والمستحقات",
     "source": "فروق ميزانيتين متوازنتين + قائمة الدخل (لا يُعتمد شيت CASH FLOW في الملف)",
     "period": "بين تاريخي ميزانية"},
    {"kpi": "حركة جاري الشركاء/المسحوبات", "formula": "التغير في حقوق الملكية − صافي ربح الفترة",
     "source": "B.S. كود 41 + قائمة الدخل", "period": "بين تاريخي ميزانية"},
]


def build() -> dict:
    import openpyxl  # local import: keeps the module importable without it
    wb_pop = openpyxl.load_workbook(F_POP, data_only=True)
    bs, flags = parse_balance_sheets(wb_pop)
    pop_is = parse_pop_is(wb_pop)
    oca = parse_other_current_assets(wb_pop)
    months, is_flags = parse_income_detail(openpyxl.load_workbook(F_IS_DETAIL, data_only=True))
    months = enrich_income(months, pop_is)
    cfs = cash_flows(bs, months)
    as_of = max(bs)
    wc = working_capital(bs, months, as_of)
    recon = sales_reconciliation(months)
    flags = flags + is_flags + quality_flags(bs, months, pop_is, oca, recon)
    loans = [{"name": n, "kind": kind, "balances": {d: b[key] for d, b in bs.items()}}
             for n, kind, key in (("قرض المبادرة", "طويل الأجل بفائدة مدعومة", "initiative_loan"),
                                  ("قرض السيارة", "طويل الأجل", "car_loan"),
                                  ("قروض الشركاء", "قصير الأجل", "shareholder_loans"))]
    loans_feb = [i for i in oca.get("2026-02-28", []) if i["category"] in ("قروض", "قروض قصيرة")]
    return {
        "meta": {"generated": date.today().isoformat(), "as_of": as_of,
                 "balance_sheet_dates": sorted(bs), "income_months": list(months),
                 "sources": [F_POP.name, F_IS_DETAIL.name],
                 "excluded_columns_note": "عمود 28/2/2026 مستبعد لعدم توازنه"},
        "balance_sheets": bs, "bs_labels": BS_LABELS_AR,
        "income": months, "income_pop_sheet": pop_is,
        "cash_flows": cfs, "working_capital": wc,
        "other_current_assets": oca, "loans": loans, "loans_feb_2026_detail": loans_feb,
        "sales_reconciliation": recon, "quality_flags": flags,
        "alerts": alerts(bs, months, wc, cfs), "alert_thresholds": ALERTS,
        "budget": load_budget(), "kpi_dictionary": KPI_DICTIONARY,
    }


def write(payload: dict) -> None:
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    OUT_JS.write_text("window.DASH_FS = " + json.dumps(payload, ensure_ascii=False) + ";\n", encoding="utf-8")


def main() -> int:
    p = build()
    write(p)
    b = p["balance_sheets"][p["meta"]["as_of"]]
    print(f"● Statements: BS dates {p['meta']['balance_sheet_dates']} · months {p['meta']['income_months']}")
    print(f"  total assets {b['total_assets']:,.0f} · equity {b['total_equity']:,.0f} · net debt {b['net_debt']:,.0f}")
    for cf in p["cash_flows"]:
        print(f"  CF {cf['from']}→{cf['to']}: CFO {cf['cfo']:,.0f} CFI {cf['cfi']:,.0f} CFF {cf['cff']:,.0f} "
              f"check {cf['check_diff']:,.2f}")
    print(f"  {len(p['quality_flags'])} quality flags · {len(p['alerts'])} alerts")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
